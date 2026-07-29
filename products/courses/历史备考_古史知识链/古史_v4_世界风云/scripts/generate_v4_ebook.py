# -*- coding: utf-8 -*-
"""
古史_v4_世界风云  电子书生成器
============================================================
从 6集 episode Python 数据文件生成 MD + DOCX 电子书。

输出：
  - 世界风云_电子书.md
  - 世界风云_电子书.docx
"""

import os, sys, json
from datetime import date

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))

def load_episodes():
    episodes = []
    files = [
        "古史_v4_ep01_帝国的崩塌.py",
        "古史_v4_ep02_赤旗与铁幕的升起.py",
        "古史_v4_ep03_大十字.py",
        "古史_v4_ep04_冷战铁幕.py",
        "古史_v4_ep05_帝国的黄昏.py",
        "古史_v4_ep06_新战国时代.py",
    ]
    for fname in files:
        fpath = os.path.join(SCRIPT_DIR, fname)
        ns = {}
        with open(fpath, encoding="utf-8") as f:
            code = f.read()
        exec(compile(code, fpath, 'exec'), ns)
        episodes.append(ns['get_data']())
    return episodes

def render_episode(ep):
    md = []
    md.append(f"\n## 第{ep['ep']}集　{ep['title']}")
    md.append(f"\n*{ep['subtitle']}*\n")

    # ===== 第一幕：时间线 =====
    md.append("\n### 第一幕：时间线\n")
    tl = ep.get('timeline', {})
    if 'overview' in tl:
        md.append(f"\n{tl['overview'].strip()}\n")
    if 'events' in tl:
        md.append("\n| 年份 | 事件 | 描述 | 历史意义 |\n")
        md.append("|------|------|------|----------|\n")
        for ev in tl['events']:
            ev_desc = ev.get('description', '').replace('|', '\\|').replace('\n', ' ')
            ev_sig = ev.get('significance', '').replace('|', '\\|').replace('\n', ' ')
            md.append(f"| **{ev.get('year','')}** | **{ev.get('title','')}** | {ev_desc} | {ev_sig} |\n")

    # ===== 第二幕：核心问题 =====
    md.append("\n### 第二幕：核心问题\n")
    aq = ep.get('art_question', {})
    if 'theme' in aq:
        md.append(f"\n**{aq['theme']}**\n")
    if 'content' in aq:
        md.append(f"\n{aq['content'].strip()}\n")

    # ===== 第三幕：当代回响 =====
    md.append("\n### 第三幕：当代回响\n")
    ce = ep.get('contemporary_echo', {})
    if 'theme' in ce:
        md.append(f"\n**{ce['theme']}**\n")
    if 'content' in ce:
        md.append(f"\n{ce['content'].strip()}\n")

    # ===== 第四幕：春秋笔法 =====
    md.append("\n### 第四幕：春秋笔法\n")
    sa = ep.get('spring_autumn', {})
    if 'theme' in sa:
        md.append(f"\n**{sa['theme']}**\n")
    if 'content' in sa:
        md.append(f"\n{sa['content'].strip()}\n")

    # ===== 知识链索引 =====
    chains = ep.get('knowledge_chains', [])
    if chains:
        md.append("\n### 附录：本集知识链\n")
        for chain in chains:
            md.append(f"\n**{chain['chain_id']} · {chain['chain_title']}**（{chain.get('chain_type','')}）\n")
            for node in chain.get('key_nodes', []):
                weight = node.get('exam_weight', '')
                md.append(f"- **{node['node_title']}** {weight}：{node.get('content','')[:60]}…\n")
    return '\n'.join(md)

def render_chain_index(episodes):
    md = ["\n\n---\n\n## 知识链总索引\n"]
    all_chains = []
    for ep in episodes:
        for chain in ep.get('knowledge_chains', []):
            all_chains.append((chain, ep))
    for i, (chain, ep) in enumerate(all_chains, 1):
        md.append(f"\n**{i}. {chain['chain_id']} · {chain['chain_title']}**\n")
        md.append(f"来源：第{ep['ep']}集《{ep['title']}》\n")
        md.append(f"类型：{chain.get('chain_type','')}  |  节点数：{len(chain.get('key_nodes',[]))}\n")
    return '\n'.join(md)

def build_markdown(episodes):
    today = date.today().isoformat()
    lines = [
        "# 世界风云：1914-2020",
        "### 从帝国的崩塌到新战国时代",
        f"\n*生成日期：{today} | 共{len(episodes)}集*\n",
        "\n---\n",
        "\n## 导论：世界是中国的舞台",
        "\n> 本卷讲述1914-2020年的世界历史。",
        "从一战的废墟到冷战的终结，从殖民帝国的瓦解到中国的重新崛起，",
        "这106年世界历史，是中国备考历史科目最重要的外部背景。",
        "不了解这个世界，就无法真正理解中国何以走到今天。",
        "\n---\n",
    ]
    for ep in episodes:
        lines.append(render_episode(ep))
    lines.append(render_chain_index(episodes))
    return '\n'.join(lines)

def to_docx(md_text, output_path):
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        doc = Document()
        # 处理简单Markdown
        for line in md_text.split('\n'):
            line = line.strip()
            if not line:
                doc.add_paragraph()
            elif line.startswith('# '):
                p = doc.add_heading(line[2:], level=1)
            elif line.startswith('## '):
                p = doc.add_heading(line[3:], level=2)
            elif line.startswith('### '):
                p = doc.add_heading(line[4:], level=3)
            elif line.startswith('> '):
                p = doc.add_paragraph(line[2:])
                p.runs[0].italic = True
            elif line.startswith('- '):
                doc.add_paragraph(line[2:], style='List Bullet')
            elif line.startswith('*') and line.endswith('*'):
                p = doc.add_paragraph(line.strip('*'))
                p.runs[0].italic = True
            else:
                # 处理表格行
                if line.startswith('|'):
                    continue  # 简化：跳过表格
                else:
                    doc.add_paragraph(line)
        doc.save(output_path)
        return True
    except ImportError:
        return False

def main():
    print("=" * 50)
    print("古史_v4_世界风云  电子书生成器")
    print("=" * 50)

    episodes = load_episodes()
    print(f"已加载 {len(episodes)} 集")

    # 生成 MD
    md = build_markdown(episodes)
    md_path = os.path.join(SCRIPT_DIR, "..", "..", "ebooks", "世界风云_电子书.md")
    os.makedirs(os.path.dirname(md_path), exist_ok=True)
    with open(md_path, 'w', encoding='utf-8') as f:
        f.write(md)
    size_kb = os.path.getsize(md_path) // 1024
    print(f"[MD] 写入成功：{md_path}（{size_kb} KB）")

    # 生成 DOCX
    docx_path = md_path.replace('.md', '.docx')
    if to_docx(md, docx_path):
        docx_kb = os.path.getsize(docx_path) // 1024
        print(f"[DOCX] 写入成功：{docx_path}（{docx_kb} KB）")
    else:
        print("[WARN] python-docx 未安装，跳过 DOCX 生成")
        print("       安装命令：pip install python-docx")

    print(f"\n[OK] 世界风云电子书生成完成！共 {len(episodes)} 集")
    print(f"     MD: {md_path}")

if __name__ == '__main__':
    main()
